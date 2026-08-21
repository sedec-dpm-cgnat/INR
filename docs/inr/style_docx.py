# -*- coding: utf-8 -*-
"""Aplica a identidade visual do INR ao DOCX produzido pelo Pandoc.

O conteúdo continua sendo gerado a partir do HTML; este passo só transforma
estilos, tabelas, tipografia e elementos de página em equivalentes nativos do
Word, mantendo o arquivo editável.
"""
import os
import sys
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


NAVY = "18244D"
INK = "1F2A4D"
INK2 = "46567D"
MUTED = "657391"
GOLD = "B35B00"
GOLD_LIGHT = "F5EBDC"
CREAM = "F7F3EA"
BORDER = "D9CCB2"
TEAL = "127C78"
WHITE = "FFFFFF"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_font(run, name, size, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(element, fill):
    props = element.get_or_add_pPr() if element.tag.endswith("p") else element.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_border(props, edge, color, size="8", space="0", val="single"):
    borders = props.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        props.append(borders)
    tag = qn(f"w:{edge}")
    node = borders.find(tag)
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), val)
    node.set(qn("w:sz"), size)
    node.set(qn("w:space"), space)
    node.set(qn("w:color"), color)


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, fractions):
    total = 10040  # A4 with 16.5 mm side margins: 7.0 in usable width.
    widths = [int(total * f) for f in fractions]
    widths[-1] += total - sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths) - 1)])


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_para_border(paragraph, edge="bottom", color=BORDER, size="8"):
    set_border(paragraph._p.get_or_add_pPr(), edge, color, size)


def set_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    set_para_border(p, "top", BORDER, "6")
    r = p.add_run("Índice Nacional de Risco  ·  Documento executivo  ·  ")
    set_font(r, "Arial", 8, MUTED)
    set_page_field(p)
    for run in p.runs[1:]:
        set_font(run, "Arial", 8, MUTED)


def separate_adjacent_tables(doc):
    """Insert a real paragraph between adjacent tables.

    Pandoc can emit consecutive tables without a paragraph between them. Word
    and LibreOffice may then carry the first table's repeating header into the
    second table. A tiny spacer makes the table boundaries unambiguous.
    """
    body = doc._element.body
    children = list(body)
    for i in range(len(children) - 2, -1, -1):
        current = children[i]
        following = children[i + 1]
        if current.tag.endswith("tbl") and following.tag.endswith("tbl"):
            spacer = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "60")
            ppr.append(spacing)
            spacer.append(ppr)
            current.addnext(spacer)


def remove_redundant_trailing_metadata(doc):
    """Drop Pandoc's unstyled duplicate metadata tail.

    The same metadata is already present in the branded opening block. Pandoc
    emits these three labels as separate trailing paragraphs; in LibreOffice
    they can create an otherwise blank final page.
    """
    prefixes = ("ElaboraçãoCoordenação", "InstrumentosDois", "VersãoDocumento")
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith(prefixes):
            paragraph._element.getparent().remove(paragraph._element)


def style_paragraph(paragraph, font="Georgia", size=10.5, color=INK, after=6, line=1.12):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    for run in paragraph.runs:
        set_font(run, font, size, color)


def is_callout(paragraph):
    if not paragraph.runs:
        return False
    first = paragraph.runs[0].text.strip()
    if not first or not paragraph.runs[0].bold:
        return False
    return first.upper() == first and len(first) > 8 and ("—" in first or ":" in first)


def style_callout(paragraph):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(7)
    pf.space_after = Pt(9)
    pf.left_indent = Mm(5)
    pf.right_indent = Mm(3)
    pf.line_spacing = 1.12
    shade(paragraph._p, CREAM)
    set_border(paragraph._p.get_or_add_pPr(), "left", GOLD, "18")
    set_border(paragraph._p.get_or_add_pPr(), "top", BORDER, "4")
    set_border(paragraph._p.get_or_add_pPr(), "bottom", BORDER, "4")
    for i, run in enumerate(paragraph.runs):
        set_font(run, "Arial" if i == 0 else "Georgia", 10.2, GOLD if i == 0 else INK, bold=(i == 0) or run.bold, italic=run.italic)


def table_fractions(cols, first_text):
    text = first_text.lower()
    if cols == 2:
        return [0.72, 0.28] if text.startswith("item") else [0.5, 0.5]
    if cols == 3:
        if any(text.startswith(x) for x in ("patologia", "laboratório", "unidade", "resultado", "dimensão")):
            return [0.25, 0.35, 0.40]
        return [0.24, 0.32, 0.44]
    if cols == 4:
        return [0.28, 0.39, 0.18, 0.15] if text.startswith("rubrica") else [0.24, 0.25, 0.25, 0.26]
    if cols == 5:
        return [0.15, 0.39, 0.12, 0.14, 0.20]
    if cols == 6:
        if text.startswith("rubrica"):
            return [0.34, 0.14, 0.12, 0.12, 0.12, 0.16]
        return [0.40, 0.22, 0.08, 0.08, 0.09, 0.13]
    return [1.0 / cols] * cols


def looks_like_header(text):
    key = " ".join(text.lower().split())
    starts = (
        "instrumento", "patologia", "subgrupo cobrade", "laboratório", "unidade / programa",
        "função no projeto", "rubrica", "item", "semestre", "entrega", "resultado", "dimensão", "meta",
    )
    return key.startswith(starts)


def style_table(table, index):
    if not table.rows:
        return
    first_text = table.cell(0, 0).text.strip()
    set_table_geometry(table, table_fractions(len(table.columns), first_text))
    set_table_borders(table)
    header = looks_like_header(first_text)
    if header:
        repeat_header(table.rows[0])
    for r, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            fill = NAVY if header and r == 0 else (GOLD_LIGHT if (r + (0 if header else 1)) % 2 == 0 else WHITE)
            if index in (0, 2, 4):
                fill = CREAM if r % 2 == 0 else WHITE
            set_cell_shading(cell, fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    if header and r == 0:
                        set_font(run, "Arial", 8.2, WHITE, bold=True)
                    else:
                        set_font(run, "Arial", 8.6, INK2, bold=run.bold, italic=run.italic)
        if header and r == 0:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        if r == len(table.rows) - 1 and not header:
            for cell in row.cells:
                set_cell_shading(cell, GOLD_LIGHT)


def main(src, dst):
    doc = Document(src)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(15)
        section.bottom_margin = Mm(15)
        section.left_margin = Mm(16.5)
        section.right_margin = Mm(16.5)
        section.header_distance = Mm(7)
        section.footer_distance = Mm(8)
        add_footer(section)

    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph"):
        st = styles[name]
        st.font.name = "Georgia"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
        st.font.size = Pt(10.5)
        st.font.color.rgb = rgb(INK)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.12
    if "Compact" in styles:
        st = styles["Compact"]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(9.4)
        st.font.color.rgb = rgb(INK2)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (("Heading 1", 19, 10, 6), ("Heading 2", 17, 13, 6), ("Heading 3", 12.5, 9, 4)):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(NAVY)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0

    # Branded opening block: navy masthead with the existing logos and metadata.
    for i in range(min(5, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        shade(p._p, NAVY)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0 if i < 4 else 8)
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.right_indent = Mm(4)
        for run in p.runs:
            set_font(run, "Arial" if i in (0, 2, 4) else "Georgia", {0: 25, 2: 8.2, 3: 11, 4: 8.5, 1: 8}[i], WHITE, bold=(i in (0, 2, 4)) or run.bold, italic=run.italic)
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(5)
        elif i == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(5)
        elif i == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(7)
        elif i == 4:
            set_para_border(p, "bottom", GOLD, "18")

    for p in doc.paragraphs[5:]:
        if p.style.name.startswith("Heading"):
            for run in p.runs:
                set_font(run, "Arial", 17 if p.style.name == "Heading 2" else 12.5, NAVY, bold=True, italic=run.italic)
            if p.style.name == "Heading 2":
                set_para_border(p, "bottom", BORDER, "8")
        elif p.style.name == "Compact":
            for run in p.runs:
                set_font(run, "Arial", 9.4, INK2, bold=run.bold, italic=run.italic)
        else:
            style_paragraph(p)
            if is_callout(p):
                style_callout(p)

    for i, table in enumerate(doc.tables):
        style_table(table, i)

    separate_adjacent_tables(doc)
    remove_redundant_trailing_metadata(doc)

    # Make the generated title block and section transitions stay together.
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True

    doc.core_properties.title = "Índice Nacional de Risco — Documento Executivo"
    doc.core_properties.subject = "Programa INR em duas fases"
    doc.core_properties.author = "SEDEC/MIDR — DPM/CGNAT"
    doc.save(dst)
    print(f"DOCX estilizado: {dst} ({os.path.getsize(dst) / 1024:.0f} KB)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("uso: style_docx.py entrada.docx saida.docx")
    main(sys.argv[1], sys.argv[2])
